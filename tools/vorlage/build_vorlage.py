# -*- coding: utf-8 -*-
"""
Baut die Vorlage  Finanzen_Vorlage.xlsm  (Blätter mit openpyxl, VBA mit ovba.py).

    python3 tools/vorlage/build_vorlage.py            -> Finanzen_Vorlage.xlsm + VBA-Module-Vorlage.zip im Repo-Stamm
    python3 tools/vorlage/build_vorlage.py ziel.xlsm  -> nur die Mappe an diesem Pfad
"""
import io
import os
import re
import sys
import zipfile

import openpyxl
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import ovba  # noqa: E402

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.abspath(os.path.join(HERE, "..", ".."))

# ----------------------------------------------------------------------------- Farben / Stile
BLAU_DUNKEL = "1F3864"
BLAU_BUTTON = "2E75B6"
BLAU_HELL = "D9E1F2"
BOX = "EEF2F8"
GRAU = "595959"
GELB = "FFFF99"
GRUEN_HELL = "E2EFDA"
BLAU_ZART = "DDEBF7"
GELB_ZART = "FFF2CC"
ORANGE_ZART = "FCE4D6"
EURO = '#,##0.00" €";[Red]\\-#,##0.00" €";\\–'
EURO_PLAIN = '#,##0.00" €"'
DATUM = "DD.MM.YYYY"
MONAT = "MMM YYYY"

F_TITEL = Font(name="Calibri", size=16, bold=True, color=BLAU_DUNKEL)
F_NOTE = Font(name="Calibri", size=9, color=GRAU)
F_NOTE8 = Font(name="Calibri", size=8, color=GRAU)
F_KOPF = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
F_FETT = Font(name="Calibri", size=10, bold=True)
F_NORM = Font(name="Calibri", size=10)
F_KLEIN = Font(name="Calibri", size=9, color="8C8C8C")
F_EINGABE = Font(name="Calibri", size=10, color="0000FF")
F_BUTTON = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
F_KPI = Font(name="Calibri", size=16, bold=True, color=BLAU_DUNKEL)
F_KPILABEL = Font(name="Calibri", size=9, bold=True, color=GRAU)
FILL_KOPF = PatternFill("solid", fgColor=BLAU_DUNKEL)
FILL_SEKTION = PatternFill("solid", fgColor=BLAU_HELL)
FILL_BOX = PatternFill("solid", fgColor=BOX)
FILL_BUTTON = PatternFill("solid", fgColor=BLAU_BUTTON)
FILL_GELB = PatternFill("solid", fgColor=GELB)
THIN = Side(style="thin", color="BFBFBF")
RAHMEN = Border(bottom=THIN)


def titel(ws, text, note=None):
    ws["A1"] = text
    ws["A1"].font = F_TITEL
    if note:
        ws["A2"] = note
        ws["A2"].font = F_NOTE
    ws.sheet_view.showGridLines = False


def kopf(ws, row, labels, start_col=1):
    for i, lab in enumerate(labels):
        c = ws.cell(row=row, column=start_col + i, value=lab)
        c.font = F_KOPF
        c.fill = FILL_KOPF
        c.alignment = Alignment(vertical="center", wrap_text=True)
    ws.row_dimensions[row].height = 28


def sektion(ws, coord, text, bis=None):
    ws[coord] = text
    ws[coord].font = F_FETT
    ws[coord].fill = FILL_SEKTION
    if bis:
        ws.merge_cells(f"{coord}:{bis}")


def button(ws, coord, text, bis=None):
    c = ws[coord]
    c.value = text
    c.font = F_BUTTON
    c.fill = FILL_BUTTON
    c.alignment = Alignment(horizontal="center", vertical="center")
    c.hyperlink = f"#'{ws.title}'!{coord}"
    if bis:
        ws.merge_cells(f"{coord}:{bis}")
    ws.row_dimensions[c.row].height = 22


def kpi(ws, col, row, label, formel, fmt, note, ncols=2):
    c0 = get_column_letter(col)
    c1 = get_column_letter(col + ncols - 1)
    for r in (row, row + 1, row + 2):
        ws.merge_cells(f"{c0}{r}:{c1}{r}")
        for cc in range(col, col + ncols):
            ws.cell(row=r, column=cc).fill = FILL_BOX
    ws[f"{c0}{row}"] = label
    ws[f"{c0}{row}"].font = F_KPILABEL
    ws[f"{c0}{row + 1}"] = formel
    ws[f"{c0}{row + 1}"].font = F_KPI
    ws[f"{c0}{row + 1}"].number_format = fmt
    ws[f"{c0}{row + 2}"] = note
    ws[f"{c0}{row + 2}"].font = F_NOTE8
    ws.row_dimensions[row + 1].height = 24


def widths(ws, w):
    for col, val in w.items():
        ws.column_dimensions[col].width = val


# ----------------------------------------------------------------------------- Blätter
AKTIV = 'Abos!$C$6:$C$205,"<>beendet",Abos!$C$6:$C$205,"<>abgeschlossen",Abos!$C$6:$C$205,"<>",Abos!$S$6:$S$205,""'


def blatt_uebersicht(wb):
    ws = wb.active
    ws.title = "Übersicht"
    ws.sheet_properties.codeName = "shUebersicht"
    ws.sheet_properties.tabColor = BLAU_DUNKEL
    titel(ws, "Finanzen – Übersicht",
          "Alles kommt aus dem George-Export (JSON): Konten, Buchungen, Abos. Blau = Eingabe · Schwarz = Formel. Import über die Schaltflächen oder das Menüband „Finanzen“.")
    widths(ws, {"A": 26, "B": 15, "C": 22, "D": 16, "E": 14, "F": 14, "G": 14, "H": 16, "I": 16, "J": 14, "K": 16, "L": 14})
    button(ws, "A4", "▶  Dateien importieren …", "B4")
    button(ws, "D4", "▶  Import-Ordner einlesen", "E4")
    button(ws, "G4", "Import-Ordner öffnen", "H4")
    button(ws, "J4", "⟳  Abos neu erkennen", "K4")
    ws["A5"] = "Letzter Import:"
    ws["A5"].font = F_NOTE
    ws["B5"].font = F_NOTE
    ws["G5"] = "Datenstand (jüngste Buchung):"
    ws["G5"].font = F_NOTE
    ws["I5"] = '=IF(COUNT(Buchungen!$A$6:$A$20000)=0,"–",MAX(Buchungen!$A$6:$A$20000))'
    ws["I5"].font = F_NOTE
    ws["I5"].number_format = DATUM
    ws["A7"] = "Auswertungsmonat"
    ws["A7"].font = F_FETT
    ws["B7"] = "=DATE(YEAR(TODAY()),MONTH(TODAY()),1)"
    ws["B7"].font = Font(name="Calibri", size=10, bold=True, color="0000FF")
    ws["B7"].fill = FILL_GELB
    ws["B7"].number_format = MONAT
    ws["C7"] = "← Datum eintragen (z. B. 1.9.2026), um einen anderen Monat auszuwerten"
    ws["C7"].font = F_NOTE

    # KPI-Reihe 1: regelmäßige Zahlungen
    kpi(ws, 1, 9, "Aktive Abos", f'=COUNTIFS(Abos!$B$6:$B$205,"Abo*",{AKTIV})', "0", "Blatt Abos, Art „Abo“")
    kpi(ws, 3, 9, "Abo-Kosten / Monat", f'=SUMIFS(Abos!$G$6:$G$205,Abos!$B$6:$B$205,"Abo*",{AKTIV})', EURO, "auf den Monat umgerechnet")
    kpi(ws, 5, 9, "Teilzahlungen laufend", f'=COUNTIFS(Abos!$B$6:$B$205,"Teilzahlung",{AKTIV})', "0", "Raten, Finanzierungen")
    kpi(ws, 7, 9, "Teilzahlungen / Monat", f'=SUMIFS(Abos!$G$6:$G$205,Abos!$B$6:$B$205,"Teilzahlung",{AKTIV})', EURO, "Summe der laufenden Raten")
    kpi(ws, 9, 9, "Daueraufträge / Monat", f'=SUMIFS(Abos!$G$6:$G$205,Abos!$B$6:$B$205,"Dauerauftrag",{AKTIV})', EURO, "Miete u. ä.")
    kpi(ws, 11, 9, "Fixe Belastung / Monat", "=C10+G10+I10", EURO, "Abos + Teilzahlungen + Daueraufträge")
    # KPI-Reihe 2: Auswertungsmonat
    B = "Buchungen!$E$6:$E$20000"
    H = "Buchungen!$H$6:$H$20000"
    J = "Buchungen!$J$6:$J$20000"
    kpi(ws, 1, 13, "Fixe Belastung / Jahr", "=K10*12", EURO, "hochgerechnet")
    kpi(ws, 3, 13, "Einnahmen im Monat", f'=SUMIFS({B},{H},"Einnahme",{J},$B$7)', EURO, "Art „Einnahme“")
    kpi(ws, 5, 13, "Ausgaben im Monat", f'=-SUMIFS({B},{H},"Ausgabe",{J},$B$7)+G14', EURO, "alle Ausgaben inkl. regelmäßige")
    kpi(ws, 7, 13, "davon regelmäßig", f'=-SUMIFS({B},{H},"Abo*",{J},$B$7)-SUMIFS({B},{H},"Teilzahlung",{J},$B$7)-SUMIFS({B},{H},"Dauerauftrag",{J},$B$7)', EURO, "Abos, Teilzahlungen, Daueraufträge")
    kpi(ws, 9, 13, "Saldo im Monat", "=C14-E14", EURO, "Einnahmen − Ausgaben")
    kpi(ws, 11, 13, "Buchungen gesamt", "=COUNT(Buchungen!$A$6:$A$20000)", "0", "alle Konten")
    ws.conditional_formatting.add("I14", CellIsRule(operator="lessThan", formula=["0"], font=Font(color="C00000", bold=True, size=16)))

    # Konten
    sektion(ws, "A16", "Konten – werden beim Import aus dem George-Export übernommen (Kontobezeichnung, IBAN/Karte); Typ ggf. anpassen", "I16")
    kopf(ws, 17, ["Konto", "Typ", "IBAN / Karte", "Buchungen", "Erste Buchung", "Letzte Buchung", "Abos laufend", "Fixe Belastung / Monat", "Saldo im Auswertungsmonat"])
    for r in range(18, 26):
        ws.cell(row=r, column=1).font = F_EINGABE
        ws.cell(row=r, column=2).font = F_EINGABE
        ws.cell(row=r, column=3).font = F_EINGABE
        ws[f"D{r}"] = f'=IF($A{r}="","",COUNTIF(Buchungen!$B$6:$B$20000,$A{r}))'
        ws[f"E{r}"] = f'=IF($A{r}="","",IFERROR(1/(1/_xlfn.MINIFS(Buchungen!$A$6:$A$20000,Buchungen!$B$6:$B$20000,$A{r})),"–"))'
        ws[f"F{r}"] = f'=IF($A{r}="","",IFERROR(1/(1/_xlfn.MAXIFS(Buchungen!$A$6:$A$20000,Buchungen!$B$6:$B$20000,$A{r})),"–"))'
        ws[f"G{r}"] = f'=IF($A{r}="","",COUNTIFS(Abos!$M$6:$M$205,$A{r},{AKTIV}))'
        ws[f"H{r}"] = f'=IF($A{r}="","",SUMIFS(Abos!$G$6:$G$205,Abos!$M$6:$M$205,$A{r},{AKTIV}))'
        ws[f"I{r}"] = f'=IF($A{r}="","",SUMIFS(Buchungen!$E$6:$E$20000,Buchungen!$B$6:$B$20000,$A{r},Buchungen!$J$6:$J$20000,$B$7))'
        for col in "DEFGHI":
            ws[f"{col}{r}"].font = F_NORM
        ws[f"D{r}"].number_format = "0"
        ws[f"E{r}"].number_format = DATUM
        ws[f"F{r}"].number_format = DATUM
        ws[f"G{r}"].number_format = "0"
        ws[f"H{r}"].number_format = EURO
        ws[f"I{r}"].number_format = EURO
    dv = DataValidation(type="list", formula1="=Einstellungen!$F$5:$F$9", allow_blank=True)
    ws.add_data_validation(dv)
    dv.add("B18:B25")

    # Fälligkeiten
    sektion(ws, "A28", "Was demnächst abgebucht wird (laufende Abos, Teilzahlungen, Daueraufträge – ohne „x“ in Ignorieren)", "I28")
    kopf(ws, 29, ["Zeitraum", "Zahlungen", "Summe"])
    L = "Abos!$L$6:$L$205"
    E = "Abos!$E$6:$E$205"
    zeilen = [
        ("nächste 7 Tage", f'">="&TODAY()', f'"<="&TODAY()+7'),
        ("nächste 30 Tage", f'">="&TODAY()', f'"<="&TODAY()+30'),
        ("überfällig (erwartet, aber noch nicht gebucht)", f'"<"&TODAY()', f'"<>"'),
    ]
    for i, (lab, c1, c2) in enumerate(zeilen):
        r = 30 + i
        ws[f"A{r}"] = lab
        ws[f"A{r}"].font = F_NORM
        ws[f"B{r}"] = f'=COUNTIFS({L},{c1},{L},{c2},{AKTIV})'
        ws[f"C{r}"] = f'=SUMIFS({E},{L},{c1},{L},{c2},{AKTIV})'
        ws[f"B{r}"].font = F_NORM
        ws[f"C{r}"].font = F_NORM
        ws[f"B{r}"].number_format = "0"
        ws[f"C{r}"].number_format = EURO
    ws["A34"] = "Hinweis: Der Datenstand ist die jüngste Buchung je Konto. Exportiere regelmäßig alle Konten, damit Abos nicht fälschlich als „überfällig“ oder „beendet“ erscheinen."
    ws["A34"].font = F_NOTE
    return ws


BUCH_KOPF = ["Datum", "Konto", "Empfänger / Auftraggeber", "Verwendungszweck / Zahlungsreferenz", "Betrag", "Zahlungsart",
             "Abo / Teilzahlung", "Art", "Notiz", "Monat", "IBAN Partner", "Gläubiger-ID", "Mandatsreferenz", "Schlüssel"]


def blatt_buchungen(wb):
    ws = wb.create_sheet("Buchungen")
    ws.sheet_properties.codeName = "shBuchungen"
    titel(ws, "Buchungen (alle Konten)")
    ws["A2"] = "Wird vom Import befüllt (Zahlungsart aus den Feldern des Exports: Mandat → Lastschrift, IBAN → Überweisung, Kreditkarte → Karte). Ändern darfst du nur die Notiz."
    ws["A2"].font = F_NOTE
    ws["A3"] = "Spalten „Abo / Teilzahlung“ und „Art“ setzt die Abo-Erkennung: Abo · Abo (variabel) · Teilzahlung · Dauerauftrag · Ausgabe · Einnahme · Umbuchung (eigene Konten, Kartenabrechnung) · Ignorieren (Betrag 0)."
    ws["A3"].font = F_NOTE
    button(ws, "D1", "▶  Dateien importieren …")
    button(ws, "F1", "▶  Import-Ordner einlesen", "G1")
    button(ws, "H1", "⟳  Abos neu erkennen", "I1")
    widths(ws, {"A": 11, "B": 14, "C": 28, "D": 44, "E": 13, "F": 12, "G": 24, "H": 13, "I": 24, "J": 10, "K": 24, "L": 20, "M": 20, "N": 8})
    kopf(ws, 5, BUCH_KOPF)
    ws.freeze_panes = "A6"
    ws.auto_filter.ref = "A5:N5005"
    for r in range(6, 3006):
        ws.cell(row=r, column=1).number_format = DATUM
        ws.cell(row=r, column=5).number_format = EURO
        ws.cell(row=r, column=10).number_format = MONAT
        ws.cell(row=r, column=10).font = F_KLEIN
        ws.cell(row=r, column=14).font = F_KLEIN
        ws.cell(row=r, column=9).font = F_EINGABE
    ws.conditional_formatting.add("A6:H5005", FormulaRule(formula=['$G6<>""'], fill=PatternFill("solid", fgColor=GRUEN_HELL)))
    ws.conditional_formatting.add("A6:N5005", FormulaRule(formula=['$H6="Umbuchung"'], font=Font(color="8C8C8C")))
    ws.conditional_formatting.add("E6:E5005", FormulaRule(formula=["$E6>0"], font=Font(color="008000")))
    ws.print_title_rows = "5:5"
    return ws


ABO_KOPF = ["Abo / Empfänger", "Art", "Status", "Rhythmus", "Betrag (zuletzt)", "Ø Betrag", "pro Monat", "pro Jahr", "Buchungen",
            "Erste Buchung", "Letzte Buchung", "Nächste (erwartet)", "Konto", "Zahlungsart", "Erkannt durch", "Raten",
            "Verwendungszweck (zuletzt)", "Kategorie", "Ignorieren", "Schlüssel"]


def blatt_abos(wb):
    ws = wb.create_sheet("Abos")
    ws.sheet_properties.codeName = "shAbos"
    ws.sheet_properties.tabColor = "548235"
    titel(ws, "Abos, Teilzahlungen und Daueraufträge – erkannt aus den Buchungen")
    ws["A2"] = ("Erkannt durch: gleicher Abbucher (Mandat / Gläubiger-ID / IBAN / Name) in festem Rhythmus · gleicher Betrag jeden Monat (auch bei wechselndem Empfängertext). "
                "Blau = Eingabe: Kategorie und „x“ bei Ignorieren bleiben bei der nächsten Erkennung erhalten.")
    ws["A2"].font = F_NOTE
    ws["A3"] = "Noch keine Erkennung gelaufen – Buchungen importieren oder „Abos neu erkennen“ klicken."
    ws["A3"].font = F_NOTE
    button(ws, "H1", "⟳  Abos neu erkennen", "I1")
    widths(ws, {"A": 30, "B": 14, "C": 13, "D": 14, "E": 13, "F": 12, "G": 12, "H": 12, "I": 10, "J": 12, "K": 12, "L": 13,
                "M": 14, "N": 12, "O": 30, "P": 10, "Q": 36, "R": 18, "S": 10, "T": 8})
    kopf(ws, 5, ABO_KOPF)
    ws.freeze_panes = "B6"
    ws.auto_filter.ref = "A5:T205"
    for r in range(6, 206):
        for c in range(1, 21):
            ws.cell(row=r, column=c).font = F_NORM
        for c in (5, 6, 7, 8):
            ws.cell(row=r, column=c).number_format = EURO_PLAIN
        for c in (10, 11, 12):
            ws.cell(row=r, column=c).number_format = DATUM
        ws.cell(row=r, column=9).number_format = "0"
        ws.cell(row=r, column=18).font = F_EINGABE
        ws.cell(row=r, column=19).font = F_EINGABE
        ws.cell(row=r, column=19).alignment = Alignment(horizontal="center")
        ws.cell(row=r, column=20).font = F_KLEIN
    ws.conditional_formatting.add("A6:T205", FormulaRule(formula=['OR($C6="beendet",$C6="abgeschlossen")'], font=Font(color="8C8C8C")))
    ws.conditional_formatting.add("A6:T205", FormulaRule(formula=['$S6<>""'], font=Font(color="BFBFBF", strike=True)))
    ws.conditional_formatting.add("C6:C205", FormulaRule(formula=['$C6="überfällig"'], fill=PatternFill("solid", fgColor=ORANGE_ZART)))
    ws.conditional_formatting.add("B6:B205", FormulaRule(formula=['LEFT($B6,3)="Abo"'], fill=PatternFill("solid", fgColor=GRUEN_HELL)))
    ws.conditional_formatting.add("B6:B205", FormulaRule(formula=['$B6="Teilzahlung"'], fill=PatternFill("solid", fgColor=BLAU_ZART)))
    ws.conditional_formatting.add("B6:B205", FormulaRule(formula=['$B6="Dauerauftrag"'], fill=PatternFill("solid", fgColor=GELB_ZART)))
    dv1 = DataValidation(type="list", formula1='"x"', allow_blank=True)
    dv2 = DataValidation(type="list", formula1="=Einstellungen!$D$5:$D$30", allow_blank=True)
    ws.add_data_validation(dv1)
    ws.add_data_validation(dv2)
    dv1.add("S6:S205")
    dv2.add("R6:R205")
    ws.print_title_rows = "5:5"
    return ws


def blatt_monat(wb):
    ws = wb.create_sheet("Monatsübersicht")
    ws.sheet_properties.codeName = "shMonat"
    titel(ws, "Monatsübersicht", "Je Monat mit Buchungen (wird von der Abo-Erkennung befüllt, die jüngsten 36 Monate). Beträge aus dem Blatt Buchungen nach Art.")
    widths(ws, {"A": 12, "B": 14, "C": 14, "D": 14, "E": 14, "F": 16, "G": 16, "H": 14})
    kopf(ws, 5, ["Monat", "Einnahmen", "Abos", "Teilzahlungen", "Daueraufträge", "sonstige Ausgaben", "Ausgaben gesamt", "Saldo"])
    B = "Buchungen!$E$6:$E$20000"
    H = "Buchungen!$H$6:$H$20000"
    J = "Buchungen!$J$6:$J$20000"
    for r in range(6, 42):
        ws[f"A{r}"].number_format = MONAT
        ws[f"A{r}"].font = F_FETT
        ws[f"B{r}"] = f'=IF($A{r}="","",SUMIFS({B},{H},"Einnahme",{J},$A{r}))'
        ws[f"C{r}"] = f'=IF($A{r}="","",-SUMIFS({B},{H},"Abo*",{J},$A{r}))'
        ws[f"D{r}"] = f'=IF($A{r}="","",-SUMIFS({B},{H},"Teilzahlung",{J},$A{r}))'
        ws[f"E{r}"] = f'=IF($A{r}="","",-SUMIFS({B},{H},"Dauerauftrag",{J},$A{r}))'
        ws[f"F{r}"] = f'=IF($A{r}="","",-SUMIFS({B},{H},"Ausgabe",{J},$A{r}))'
        ws[f"G{r}"] = f'=IF($A{r}="","",C{r}+D{r}+E{r}+F{r})'
        ws[f"H{r}"] = f'=IF($A{r}="","",B{r}-G{r})'
        for col in "BCDEFGH":
            ws[f"{col}{r}"].number_format = EURO
            ws[f"{col}{r}"].font = F_NORM
        ws[f"G{r}"].font = F_FETT
        ws[f"H{r}"].font = F_FETT
    ws.conditional_formatting.add("H6:H41", CellIsRule(operator="lessThan", formula=["0"], font=Font(color="C00000", bold=True)))
    ch = BarChart()
    ch.type = "col"
    ch.grouping = "stacked"
    ch.overlap = 100
    ch.title = "Ausgaben je Monat"
    ch.y_axis.title = "€"
    ch.height = 9
    ch.width = 26
    data = Reference(ws, min_col=3, max_col=6, min_row=5, max_row=41)
    cats = Reference(ws, min_col=1, min_row=6, max_row=41)
    ch.add_data(data, titles_from_data=True)
    ch.set_categories(cats)
    ws.add_chart(ch, "J5")
    return ws


def blatt_einstellungen(wb):
    ws = wb.create_sheet("Einstellungen")
    ws.sheet_properties.codeName = "shEinst"
    titel(ws, "Einstellungen", "Blau = änderbar. Die Abo-Erkennung läuft nach jedem Import; nach Änderungen hier „Abos neu erkennen“ klicken.")
    widths(ws, {"A": 34, "B": 44, "C": 60, "D": 22, "E": 3, "F": 16, "G": 3, "H": 16})
    kopf(ws, 4, ["Parameter", "Wert", "Erklärung"])
    params = [
        ("Import-Ordner", "", "leer = Ordner „Import“ neben der Mappe (auch in OneDrive). Eingelesene Dateien wandern nach „Import\\erledigt“."),
        ("Mindestanzahl Buchungen", 3, "So oft muss ein Abbucher bzw. ein Betrag regelmäßig vorkommen (Jahresabos: 2)."),
        ("Toleranz Betrag „fest“", 0.01, "Beträge, die höchstens so weit voneinander abweichen, gelten als gleich (sonst „Abo (variabel)“)."),
        ("Toleranz Termin (Tage)", 7, "Abweichung vom erwarteten Abstand (monatlich ≈ 30 Tage), die noch als regelmäßig gilt."),
        ("Mindestbetrag", 1, "Kleinere Beträge werden nicht als Abo betrachtet."),
        ("Stichwörter Teilzahlung", "rate;raten;ratenzahlung;ratenkauf;teilzahlung;teilbetrag;finanzierung;kreditrate;kreditrückzahlung;tilgung;leasing;abzahlung",
         "Kommt eines davon im Verwendungszweck/Empfänger vor (oder ein Zähler wie „3/12“), ist eine feste Zahlung eine „Teilzahlung“."),
        ("Stichwörter Umbuchung", "kreditkartenabrechnung;kartenabrechnung;kreditkartenrechnung;kartenausgleich;umbuchung;eigenübertrag;eigenüberweisung",
         "Solche Buchungen zählen als Umbuchung und nie als Abo (zusätzlich: Überweisungen zwischen eigenen Konten laut Übersicht)."),
        ("Stichwörter Bargeld", "bargeld;bankomat;behebung;atm;automat;auszahlung", "Behebungen sind nie Abos, auch wenn jeden Monat derselbe Betrag abgehoben wird."),
    ]
    for i, (p, v, erkl) in enumerate(params):
        r = 5 + i
        ws[f"A{r}"] = p
        ws[f"A{r}"].font = F_NORM
        ws[f"B{r}"] = v
        ws[f"B{r}"].font = F_EINGABE
        ws[f"B{r}"].fill = FILL_GELB
        ws[f"B{r}"].alignment = Alignment(wrap_text=True, vertical="top")
        ws[f"C{r}"] = erkl
        ws[f"C{r}"].font = F_NOTE
        ws[f"C{r}"].alignment = Alignment(wrap_text=True, vertical="top")
    ws["B7"].number_format = "0%"
    ws["B9"].number_format = EURO
    for r in (10, 11, 12):
        ws.row_dimensions[r].height = 42
    kopf(ws, 4, ["Kategorien (für Abos)"], start_col=4)
    kats = ["Wohnen", "Energie", "Handy & Internet", "Streaming & Abos", "Software & Cloud", "Versicherung", "Kredit & Raten", "Mobilität",
            "Fitness & Sport", "Gesundheit", "Bildung & Kinder", "Spenden & Mitgliedschaften", "Bank & Gebühren", "Sparen", "Sonstiges"]
    for i, k in enumerate(kats):
        ws.cell(row=5 + i, column=4, value=k).font = F_EINGABE
    kopf(ws, 4, ["Kontotypen"], start_col=6)
    for i, k in enumerate(["Girokonto", "Sparkonto", "Bausparkonto", "Kreditkarte", "Sonstiges"]):
        ws.cell(row=5 + i, column=6, value=k).font = F_EINGABE
    kopf(ws, 4, ["Arten (fest)"], start_col=8)
    for i, k in enumerate(["Abo", "Abo (variabel)", "Teilzahlung", "Dauerauftrag", "Ausgabe", "Einnahme", "Umbuchung", "Ignorieren"]):
        ws.cell(row=5 + i, column=8, value=k).font = F_NORM
    ws["A15"] = "Kontotypen steuern die Einordnung: Spar-/Bausparkonto → Ein-/Auszahlungen sind Umbuchungen (Zinsen = Einnahme, KESt = Ausgabe); Kreditkarte → Eingänge sind Umbuchungen (Kartenausgleich)."
    ws["A15"].font = F_NOTE
    return ws


def blatt_hinweise(wb):
    ws = wb.create_sheet("Hinweise")
    ws.sheet_properties.codeName = "shHinweise"
    titel(ws, "Hinweise")
    widths(ws, {"A": 4, "B": 130})
    bloecke = [
        ("Import", [
            "George: je Konto (Giro, Sparkonto, Kreditkarte) den Umsatz-Export als JSON herunterladen – mehr braucht die Vorlage nicht.",
            "„Dateien importieren …“ (Übersicht, Buchungen oder Menüband „Finanzen“) → alle JSON-Dateien markieren → Öffnen. Oder: Dateien in den Ordner „Import“ neben der Mappe ziehen und „Import-Ordner einlesen“ klicken.",
            "Konten werden aus dem Export angelegt (Kontobezeichnung + IBAN/Kartennummer) und in der Übersicht eingetragen. Den Typ kannst du dort anpassen.",
            "Bereits importierte Buchungen werden erkannt (Spalte „Schlüssel“) – dieselbe Datei mehrfach einlesen ist unschädlich.",
            "Nach jedem Import läuft die Abo-Erkennung automatisch über alle Buchungen.",
        ]),
        ("Wie Abos erkannt werden", [
            "Gleicher Abbucher: Buchungen mit derselben Mandatsreferenz, Gläubiger-ID, IBAN oder demselben Namen werden als Serie betrachtet. Kommen sie in festem Abstand (wöchentlich, monatlich, vierteljährlich, halbjährlich, jährlich – Toleranz in den Einstellungen) und mindestens so oft wie „Mindestanzahl“, ist es eine regelmäßige Zahlung.",
            "Fester Betrag → „Abo“ (Lastschrift/Karte), „Dauerauftrag“ (Überweisung) oder „Teilzahlung“ (Stichwort wie Rate/Ratenzahlung/Teilbetrag oder Zähler wie „3/12“ im Text). Schwankender Betrag (Handy, Strom) → „Abo (variabel)“. Preisänderungen (z. B. 10,99 → 11,99) bleiben ein Abo.",
            "Mehrere Abos beim selben Anbieter (z. B. zwei Apple-Abos am selben Tag, Jahresabo neben Einzelkäufen) werden über den Betrag getrennt.",
            "Gleicher Betrag: Zahlungen mit exakt demselben Betrag in festem Rhythmus (eine je Zeitraum) gelten auch dann als Teilzahlung/Dauerauftrag/Abo, wenn der Empfängertext wechselt.",
            "Nie Abos: Umbuchungen zwischen eigenen Konten, Kreditkartenabrechnungen, Bargeldbehebungen, Eingänge, Beträge unter dem Mindestbetrag.",
            "Status: aktiv = nächste Zahlung steht noch aus · überfällig = erwartet, aber (laut Datenstand des Kontos) nicht gekommen · beendet = seit mehr als einem Zeitraum nichts mehr · abgeschlossen = letzte Rate (Zähler) erreicht.",
            "„pro Monat“ rechnet den Rhythmus um (vierteljährlich ÷ 3, jährlich ÷ 12, wöchentlich × 52 ÷ 12). In der Übersicht zählen nur Zeilen ohne „x“ bei Ignorieren.",
        ]),
        ("Eingaben", [
            "Blatt Abos: Kategorie und „x“ bei Ignorieren bleiben bei der nächsten Erkennung erhalten (Zuordnung über die Spalte „Schlüssel“).",
            "Blatt Buchungen: nur die Notiz ist für Eingaben gedacht; Abo/Art werden von der Erkennung neu gesetzt.",
            "Einstellungen: Mindestanzahl, Toleranzen, Mindestbetrag und Stichwörter. Danach „Abos neu erkennen“.",
        ]),
        ("Technik", [
            "Die Mappe enthält Makros (.xlsm). Beim ersten Öffnen ggf. „Inhalt aktivieren“ bzw. in den Dateieigenschaften „Zulassen“ anklicken.",
            "Für Excel unter Windows gebaut (Excel 2019 / Microsoft 365). Die Formeln MINIFS/MAXIFS brauchen Excel 2019 oder neuer.",
            "Kontostände sind nicht im George-Export enthalten und werden deshalb nicht geführt.",
            "Quelltexte der Makros: Ordner tools/vorlage/vba im Repository bzw. VBA-Module-Vorlage.zip.",
        ]),
    ]
    r = 3
    for kopfz, zeilen in bloecke:
        ws[f"A{r}"] = kopfz
        ws[f"A{r}"].font = F_FETT
        ws[f"A{r}"].fill = FILL_SEKTION
        ws[f"B{r}"].fill = FILL_SEKTION
        r += 1
        for i, z in enumerate(zeilen):
            ws[f"A{r}"] = str(i + 1) if kopfz == "Import" else "•"
            ws[f"A{r}"].font = F_FETT
            ws[f"A{r}"].alignment = Alignment(vertical="top")
            ws[f"B{r}"] = z
            ws[f"B{r}"].font = F_NORM
            ws[f"B{r}"].alignment = Alignment(wrap_text=True, vertical="top")
            ws.row_dimensions[r].height = 15 * max(1, (len(z) // 120) + 1)
            r += 1
        r += 1
    return ws


# ----------------------------------------------------------------------------- VBA + xlsm
def read_src(name):
    with open(os.path.join(HERE, "vba", name), encoding="utf-8") as f:
        return f.read()


def vba_modules(wb):
    mods = [("ThisWorkbook", "workbook", read_src("ThisWorkbook.cls"))]
    for ws in wb.worksheets:
        mods.append((ws.sheet_properties.codeName, "sheet", ""))
    mods.append(("modJson", "std", read_src("modJson.bas")))
    mods.append(("modImport", "std", read_src("modImport.bas")))
    mods.append(("modAbo", "std", read_src("modAbo.bas")))
    return mods


def xlsx_to_xlsm(xlsx_bytes, vba_bin, custom_ui):
    """Macht aus der openpyxl-Mappe eine .xlsm: vbaProject.bin, Menüband, codeName der Mappe."""
    zin = zipfile.ZipFile(io.BytesIO(xlsx_bytes))
    out = io.BytesIO()
    zout = zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED)
    for item in zin.infolist():
        data = zin.read(item.filename)
        if item.filename == "[Content_Types].xml":
            s = data.decode("utf-8")
            s = s.replace("application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml",
                          "application/vnd.ms-excel.sheet.macroEnabled.main+xml")
            if 'Extension="bin"' not in s:
                s = s.replace("<Default Extension=", '<Default Extension="bin" ContentType="application/vnd.ms-office.vbaProject"/><Default Extension=', 1)
            data = s.encode("utf-8")
        elif item.filename == "xl/workbook.xml":
            s = data.decode("utf-8")
            if "codeName=" not in s:
                if "<workbookPr" in s:
                    s = s.replace("<workbookPr", '<workbookPr codeName="ThisWorkbook"', 1)
                else:
                    s = s.replace("<sheets>", '<workbookPr codeName="ThisWorkbook"/><sheets>', 1)
            data = s.encode("utf-8")
        elif item.filename == "xl/_rels/workbook.xml.rels":
            s = data.decode("utf-8")
            s = s.replace("</Relationships>",
                          '<Relationship Type="http://schemas.microsoft.com/office/2006/relationships/vbaProject" Target="vbaProject.bin" Id="rIdVBA"/></Relationships>')
            data = s.encode("utf-8")
        elif item.filename == "_rels/.rels":
            s = data.decode("utf-8")
            s = s.replace("</Relationships>",
                          '<Relationship Type="http://schemas.microsoft.com/office/2007/relationships/ui/extensibility" Target="customUI/customUI14.xml" Id="rIdUI"/></Relationships>')
            data = s.encode("utf-8")
        zout.writestr(item, data)
    zout.writestr("xl/vbaProject.bin", vba_bin)
    zout.writestr("customUI/customUI14.xml", custom_ui)
    zout.close()
    return out.getvalue()


def build(ziel=None, zip_ziel=None, seed=None):
    wb = openpyxl.Workbook()
    blatt_uebersicht(wb)
    blatt_buchungen(wb)
    blatt_abos(wb)
    blatt_monat(wb)
    blatt_einstellungen(wb)
    blatt_hinweise(wb)
    wb.active = 0
    for ws in wb.worksheets:
        ws.sheet_view.zoomScale = 100
    wb.properties.creator = "Finanzen-Excel"
    wb.properties.title = "Finanzen – Vorlage mit Abo-Erkennung"
    buf = io.BytesIO()
    wb.save(buf)
    mods = vba_modules(wb)
    vba_bin = ovba.build_vba_project(mods, seed=seed)
    with open(os.path.join(HERE, "customUI14.xml"), encoding="utf-8") as f:
        custom_ui = f.read()
    xlsm = xlsx_to_xlsm(buf.getvalue(), vba_bin, custom_ui)
    ziel = ziel or os.path.join(ROOT, "Finanzen_Vorlage.xlsm")
    with open(ziel, "wb") as f:
        f.write(xlsm)
    print("geschrieben:", ziel, len(xlsm), "Bytes")
    if zip_ziel:
        with zipfile.ZipFile(zip_ziel, "w", zipfile.ZIP_DEFLATED) as z:
            for name, kind, src in mods:
                if kind == "std":
                    z.writestr(name + ".bas", ovba._module_source(name, kind, src).encode("cp1252"))
            z.writestr("ThisWorkbook_Code.txt", ("' In den VBA-Editor (Alt+F11) unter 'DieseArbeitsmappe' einfügen\r\n" + read_src("ThisWorkbook.cls").replace("\n", "\r\n")).encode("cp1252"))
            z.writestr("customUI14.xml", custom_ui.encode("utf-8"))
        print("geschrieben:", zip_ziel)
    return ziel


if __name__ == "__main__":
    if len(sys.argv) > 1:
        build(sys.argv[1])
    else:
        build(None, os.path.join(ROOT, "VBA-Module-Vorlage.zip"))
